# -*- coding: utf-8 -*-
"""把 项目交付汇总.md 转成 Word 文档放到桌面"""
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

SRC = r"C:\Users\Lenovo\Documents\deepseek-harness\default-workspace\cet4-handbook\项目交付汇总.md"
DST = r"C:\Users\Lenovo\Desktop\单词手账-项目交付汇总.docx"

def set_font(run, name_ascii="Microsoft YaHei", name_east="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name_ascii
    r = run._element.rPr.rFonts
    r.set(qn("w:eastAsia"), name_east)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def clean(text):
    """去掉 markdown 行内标记（**加粗** / `代码`），Word 里不该出现星号和反引号"""
    return text.replace("**", "").replace("`", "")

def add_para(doc, text, size=10.5, bold=False, color=None, space_after=4, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Pt(indent)
    run = p.add_run(clean(text))
    set_font(run, size=size, bold=bold, color=color)
    return p

with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

doc = Document()
# 页边距
for sec in doc.sections:
    sec.left_margin = Pt(54)
    sec.right_margin = Pt(54)
    sec.top_margin = Pt(54)
    sec.bottom_margin = Pt(54)

i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip():
        i += 1
        continue
    if line.startswith("# "):
        add_para(doc, line[2:], size=22, bold=True, space_after=10)
    elif line.startswith("## "):
        add_para(doc, line[3:], size=16, bold=True, color=(36, 52, 71), space_after=6)
    elif line.startswith("### "):
        add_para(doc, line[4:], size=13, bold=True, color=(77, 101, 123), space_after=4)
    elif line.startswith("|") and i + 1 < len(lines) and lines[i + 1].lstrip().startswith("|"):
        # 表格块：收集连续表格行
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            block.append(lines[i])
            i += 1
        # 过滤分隔行（|---|）
        rows = []
        for b in block:
            cells = [c.strip() for c in b.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            rows.append(cells)
        if rows:
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncols)
            table.style = "Table Grid"
            for ri, r in enumerate(rows):
                for ci in range(ncols):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    txt = r[ci] if ci < len(r) else ""
                    para = cell.paragraphs[0]
                    run = para.add_run(clean(txt))
                    set_font(run, size=9.5, bold=(ri == 0), color=(36, 52, 71) if ri == 0 else None)
            add_para(doc, "", size=6, space_after=2)
        continue
    elif line.startswith("- "):
        add_para(doc, "• " + line[2:], size=10.5, indent=12, space_after=2)
    elif re.match(r"^\d+\. ", line):
        add_para(doc, line, size=10.5, indent=12, space_after=2)
    elif line.startswith("```"):
        # 代码块：收集到结束
        i += 1
        code_lines = []
        while i < len(lines) and not lines[i].startswith("```"):
            code_lines.append(lines[i])
            i += 1
        for cl in code_lines:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.left_indent = Pt(18)
            run = p.add_run(cl)
            set_font(run, name_ascii="Consolas", name_east="Consolas", size=9, color=(80, 80, 80))
    elif line == "---":
        pass
    else:
        add_para(doc, line, size=10.5, space_after=4)
    i += 1

doc.save(DST)
print("saved:", DST)

# 验证
d2 = Document(DST)
heads = [p.text for p in d2.paragraphs if p.text.startswith("#")]
print("headings:", heads[:8])
print("tables:", len(d2.tables))
print("paragraphs:", len(d2.paragraphs))
